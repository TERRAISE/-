from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

doc = Document()

# ページ設定（A4）
section = doc.sections[0]
section.page_width = Cm(21.0)
section.page_height = Cm(29.7)
section.top_margin = Cm(2.5)
section.bottom_margin = Cm(2.5)
section.left_margin = Cm(2.5)
section.right_margin = Cm(2.5)

# スタイル設定
style_normal = doc.styles['Normal']
style_normal.font.name = 'Yu Gothic'
style_normal.font.size = Pt(10.5)
style_normal.paragraph_format.line_spacing = 1.5

# ヘルパー関数
def add_heading_styled(text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)
    return h

def add_body(text, bold=False, italic=False, align=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(10.5)
    if align:
        p.alignment = align
    return p

def add_bullet(text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    if bold_prefix:
        run = p.add_run(bold_prefix)
        run.bold = True
        run.font.size = Pt(10.5)
        run = p.add_run(text)
        run.font.size = Pt(10.5)
    else:
        run = p.add_run(text)
        run.font.size = Pt(10.5)
    return p

def add_separator():
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('― ― ― ― ― ― ― ― ― ― ― ― ― ― ―')
    run.font.color.rgb = RGBColor(0xBB, 0xBB, 0xBB)
    run.font.size = Pt(10)

def add_empty():
    doc.add_paragraph()

# ========================================
# ヘッダー部分
# ========================================
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = p.add_run('2026年●月●日')
run.font.size = Pt(10)
run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = p.add_run('株式会社TERRAISE')
run.font.size = Pt(10)
run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

add_empty()

# タイトル
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('中堅製造業向けAI人材育成プログラム\n「AX Academy」始動')
run.bold = True
run.font.size = Pt(18)
run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('2ヶ月で"自走できるAX推進人材"を育成')
run.bold = True
run.font.size = Pt(14)
run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

# サブタイトル
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('研修・コンサルの繰り返しを、これで終わらせる。\n第1期モニター5名限定、完全オンライン伴走型で開講')
run.font.size = Pt(11)
run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

add_empty()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('（※ここにメインビジュアル画像を配置）')
run.font.size = Pt(9)
run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
run.italic = True
add_empty()

# ========================================
# リード文
# ========================================
add_body(
    '株式会社TERRAISE（東京都三鷹市、代表取締役：舟橋、以下「TERRAISE」）は、'
    '中堅製造業（年商50〜500億円）のDX推進担当者・経営者を対象としたAX'
    '（エーエックス：AI変革 = AI Transformation）実践スクール「AX Academy」を開講し、'
    '第1期モニター受講生を5名限定で募集開始いたします。'
)

add_body(
    '「毎年、外部研修に数百万円を投じているのに、受講後のスキルが現場に定着しない」'
    '「コンサルに依頼するたびに費用がかさみ、いつまでも自社で回せない」'
    '——こうした声は、中堅製造業のDX推進担当者から数多く寄せられています。'
)

add_body(
    '経済産業省の試算では、2030年にIT人材が最大79万人不足するとされ、'
    '特に製造業におけるOT（制御・運用技術）とAIの両方を理解する人材の不足は極めて深刻な状況です。'
    '日本のAI教育市場は3,500〜4,500億円規模（年平均成長率30〜40%）に達する見通しであり、'
    '企業の人材投資ニーズは急速に高まっています。一方で、大手企業が2〜3年前からAX推進に着手するなか、'
    '中堅企業との格差は広がり続けています。'
)

add_body(
    'AX Academyは、予知保全・品質管理・生産最適化・SCM（サプライチェーンマネジメント）に特化した'
    '2ヶ月集中・完全オンライン・伴走型の実践プログラムを通じて、外部依存から脱却し、'
    '自社内でAXを推進できる人材を育成します。通常価格100万円（税別/1名）に対し、'
    'モニター特別価格60万円/名（税別・40%OFF）にてご提供いたします。'
)

add_separator()

# ========================================
# 背景セクション
# ========================================
add_heading_styled('【背景】中堅製造業が直面する「AI推進の壁」', level=2)

add_body(
    '政府が5年間で1兆円のリスキリング予算を掲げ、日本のリスキリング市場は'
    '1.2〜1.5兆円（年平均成長率25〜30%）へと拡大しています。'
    'しかし、中堅製造業の現場には、依然として根深い4つの構造的課題が存在します。'
)

add_empty()
p = doc.add_paragraph()
run = p.add_run('1. 外部研修を繰り返すがスキルが社内に残らない')
run.bold = True
run.font.size = Pt(10.5)

add_body(
    '年間数百万円を研修費に投じても、受講後に知識やスキルが現場で活かされず、'
    '同じテーマの研修を繰り返す悪循環に陥っています。'
    '研修が「学びの場」で終わり、「実践の場」につながらないことが最大の原因です。'
)

add_empty()
p = doc.add_paragraph()
run = p.add_run('2. 外部コンサルへの「依存体質」から脱却できない')
run.bold = True
run.font.size = Pt(10.5)

add_body(
    'AI関連の案件が発生するたびに外部コンサルタントへ数百万円単位で発注する構造が常態化しています。'
    'プロジェクトが終われば知見は社外へ流出し、自社にノウハウが蓄積されません。'
    '結果として、次の案件でも再び外部に頼らざるを得ない依存体質が固定化しています。'
)

add_empty()
p = doc.add_paragraph()
run = p.add_run('3. 汎用的なAI研修では製造現場の課題を解決できない')
run.bold = True
run.font.size = Pt(10.5)

add_body(
    '市場に流通するAI研修の多くはChatGPTなど汎用ツールの活用が中心であり、'
    '予知保全や品質管理AIといった製造業固有の領域をカバーしていません。'
    'OTとAIの交差点を体系的に学べるプログラムは極めて限られているのが現状です。'
)

add_empty()
p = doc.add_paragraph()
run = p.add_run('4. 大手企業との格差が拡大し続けている')
run.bold = True
run.font.size = Pt(10.5)

add_body(
    '大手製造業はすでに2〜3年前からAXの推進体制を構築し、社内にAI人材を抱え始めています。'
    '経済産業省が試算する2030年の最大79万人のIT人材不足'
    '（経済産業省「IT人材需給に関する調査」）を前に、'
    '中堅企業にとって今この瞬間が、競争力を維持できるか否かの分岐点です。'
)

add_empty()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('（※ここに「中堅製造業の4つの構造的課題」図解を配置）')
run.font.size = Pt(9)
run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
run.italic = True

add_body(
    'AX Academyは、こうした構造的課題を抱える中堅製造業に対し、'
    '「研修でもコンサルでもない、第三の選択肢」として設計されました。'
)

add_separator()

# ========================================
# AX Academyとは
# ========================================
add_heading_styled('【AX Academyとは】製造業のための3つの約束', level=2)

add_body(
    'AX Academyは、製造業に特化したAI人材育成プログラムです。'
    'DXのさらに先にある概念「AX（AI Transformation＝AI変革）」を推進できるリーダーを、'
    '貴社の中に育てます。私たちは、受講企業の皆さまに3つの約束をお届けします。'
)

# 約束1
add_empty()
p = doc.add_paragraph()
run = p.add_run('〈約束1〉外部依存のコストサイクルを、断ち切る')
run.bold = True
run.font.size = Pt(12)
run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

add_body(
    '多くの製造業企業が、AI活用のために外部研修や外部コンサルティングを繰り返し利用しています。'
    '外部研修は「受けるたびに費用」が発生し、外部コンサルは「案件ごとに数百万円」のコストがかかります。'
    'しかも、そのたびにノウハウは社外に留まったまま。'
    'この終わりのないコストサイクルこそ、製造業のAI活用を阻む最大の構造的課題です。'
)

add_body(
    'AX Academyでは、1名への投資で社内にAX推進リーダーが定着します。'
    '外部に頼り続ける構造から脱却し、社内で自走できる体制を築くことが可能です。'
)

add_body('さらに、公的助成金の活用により、受講費用の実質負担を大幅に抑えることができます。')

add_bullet('人材開発支援助成金：受講費用の最大75%を助成')
add_bullet('教育訓練給付金：受講費用の最大70%を給付')

add_body(
    '「研修・コンサルの繰り返しを、これで終わらせる。」——それが、最初の約束です。',
    bold=True
)

# 約束2
add_empty()
p = doc.add_paragraph()
run = p.add_run('〈約束2〉製造業の現場で、即使えるAIスキルを届ける')
run.bold = True
run.font.size = Pt(12)
run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

add_body(
    '一般的なAIスクールで学ぶのは、汎用的なプログラミングやデータサイエンスの知識です。'
    'しかし、製造業の現場が本当に求めているのは、'
    '自社の設備や製品、工程に直結するAI活用スキルではないでしょうか。'
)

add_body('AX Academyのカリキュラムは、製造業の4つの重点領域に特化しています。')

add_bullet('', bold_prefix='予知保全（設備異常検知）：')
p = doc.paragraphs[-1]
run = p.add_run('突発的な設備停止を未然に防ぎ、計画的な保全を実現します')
run.font.size = Pt(10.5)

add_bullet('', bold_prefix='品質管理（画像検査AI）：')
p = doc.paragraphs[-1]
run = p.add_run('目視検査の属人化を解消し、検査精度と速度を飛躍的に向上させます')
run.font.size = Pt(10.5)

add_bullet('', bold_prefix='生産最適化：')
p = doc.paragraphs[-1]
run = p.add_run('生産計画の精度を高め、リードタイムの短縮やコスト削減を実現します')
run.font.size = Pt(10.5)

add_bullet('', bold_prefix='SCM（サプライチェーンマネジメント）：')
p = doc.paragraphs[-1]
run = p.add_run('需要予測の高度化により、在庫の最適化と欠品リスクの低減を図ります')
run.font.size = Pt(10.5)

add_body(
    '最大の特徴は、受講企業自身の現場データや業務課題を「教材」として使用する点です。'
    '架空のケーススタディではなく、実際に自社が抱える課題に取り組むからこそ、'
    '学んだスキルがそのまま現場で活きます。'
)

add_body(
    'また、AX Academyはチクセントミハイのフロー理論に基づく「フロー体験設計」を採用しています。'
    '難しすぎず、易しすぎない絶妙な課題設定により、成功体験を段階的に積み上げていく設計です。'
    '「AIは難しそう」という先入観を持つ受講者でも、着実にスキルを身につけていくことができます。'
)

# 約束3
add_empty()
p = doc.add_paragraph()
run = p.add_run('〈約束3〉2ヶ月で「動ける人材」を送り出す')
run.bold = True
run.font.size = Pt(12)
run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

add_body(
    '「学んだはずなのに、現場で何も変わらなかった」'
    '——研修にありがちなこの問題を、AX Academyは構造的に解決します。'
)

add_body(
    'AX Academyの修了基準は明確です。2ヶ月の集中期間の中で、'
    '受講者はPoC（概念実証）1件を自社の現場で実際に実装した状態で修了します。'
    '単なる座学の修了証ではなく、「現場で動かした実績」が修了の証です。'
)

add_body(
    '最終発表の場には、社内外の意思決定者の同席が可能です。'
    '受講者の成果を経営層や関係部門の責任者に直接プレゼンテーションすることで、'
    'プロジェクトの社内展開をスムーズに後押しします。'
)

add_body(
    'さらに、修了後もフォローアップ体制を用意しています。'
    'プログラム終了後に直面する実務上の課題にも、継続的なサポートで対応いたします。'
)

add_separator()

# ========================================
# カリキュラム概要
# ========================================
add_heading_styled('【カリキュラム概要】2ヶ月・3フェーズの実践プログラム', level=2)

add_body(
    'AX Academyのカリキュラムは、2ヶ月間を3つのフェーズに分けた実践型プログラムです。'
    '完全オンラインで受講可能なため、全国どこからでも参加いただけます。'
    'また、お問い合わせから最短2週間でキックオフが可能です。'
)

add_empty()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('（※ここにカリキュラム全体像のフロー図を配置）')
run.font.size = Pt(9)
run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
run.italic = True

# Phase 1
add_empty()
p = doc.add_paragraph()
run = p.add_run('〈Phase 1〉AXの全体像と思考法（Week 1〜3）')
run.bold = True
run.font.size = Pt(11)
run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

add_body('最初の3週間では、AXの基礎となる全体像と思考法を身につけます。')

add_bullet('', bold_prefix='AXとは何か：')
p = doc.paragraphs[-1]
run = p.add_run('AI Transformationの本質と、DXとの違いを理解します')
run.font.size = Pt(10.5)

add_bullet('', bold_prefix='AI活用領域マップ：')
p = doc.paragraphs[-1]
run = p.add_run('製造業においてAIが活用できる領域を体系的に把握します')
run.font.size = Pt(10.5)

add_bullet('', bold_prefix='AIプロジェクトの進め方：')
p = doc.paragraphs[-1]
run = p.add_run('構想から実装、効果検証までの一連のプロセスを学びます')
run.font.size = Pt(10.5)

add_bullet('', bold_prefix='自社AX戦略立案：')
p = doc.paragraphs[-1]
run = p.add_run('自社の課題を整理し、どの領域からAI活用を始めるべきかを設計します')
run.font.size = Pt(10.5)

# Phase 2
add_empty()
p = doc.add_paragraph()
run = p.add_run('〈Phase 2〉製造業AIの実践技術（Week 4〜7）')
run.bold = True
run.font.size = Pt(11)
run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

add_body('4週間にわたる実践フェーズでは、製造業に特化したAI技術を、手を動かしながら習得します。')

add_bullet('', bold_prefix='データ収集と前処理：')
p = doc.paragraphs[-1]
run = p.add_run('現場のデータを「AIが使える状態」に整える技術を学びます')
run.font.size = Pt(10.5)

add_bullet('', bold_prefix='予知保全AI（異常検知）：')
p = doc.paragraphs[-1]
run = p.add_run('設備センサーデータから異常の予兆を検知するモデルを構築します')
run.font.size = Pt(10.5)

add_bullet('', bold_prefix='品質管理AI（画像検査）：')
p = doc.paragraphs[-1]
run = p.add_run('製品画像から不良品を自動判定するAIを実装します')
run.font.size = Pt(10.5)

add_bullet('', bold_prefix='生産最適化と需要予測：')
p = doc.paragraphs[-1]
run = p.add_run('生産計画の精度向上や需要変動への対応力を高めるAIを開発します')
run.font.size = Pt(10.5)

add_body(
    'このフェーズでは、受講企業の実際のデータや課題を素材として使用します。'
    '汎用的なサンプルデータではなく、自社データで学ぶからこそ、'
    '修了後すぐに現場で展開できるスキルが身につきます。'
)

# Phase 3
add_empty()
p = doc.add_paragraph()
run = p.add_run('〈Phase 3〉自社課題への適用と展開（Week 8）')
run.bold = True
run.font.size = Pt(11)
run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

add_body('最終週では、Phase 1・2で培った知識と技術を、自社課題の解決に総合的に適用します。')

add_bullet('', bold_prefix='AI適用設計：')
p = doc.paragraphs[-1]
run = p.add_run('自社の優先課題に対して、最適なAIソリューションを設計します')
run.font.size = Pt(10.5)

add_bullet('', bold_prefix='社内展開計画策定：')
p = doc.paragraphs[-1]
run = p.add_run('PoCの成果を全社に広げていくためのロードマップを策定します')
run.font.size = Pt(10.5)

add_bullet('', bold_prefix='最終発表：')
p = doc.paragraphs[-1]
run = p.add_run('2ヶ月間の成果を発表します。社内外の意思決定者の同席も可能です')
run.font.size = Pt(10.5)

add_bullet('', bold_prefix='修了後フォローアップ：')
p = doc.paragraphs[-1]
run = p.add_run('プログラム修了後も、実務での課題に対してサポートを継続します')
run.font.size = Pt(10.5)

add_body(
    '全フェーズを通じて、メンターが伴走型で支援します。'
    '一方的な動画視聴や座学ではなく、受講者一人ひとりの進捗と課題に寄り添いながら進めるスタイルです。'
)

add_separator()

# ========================================
# 差別化
# ========================================
add_heading_styled('【なぜAX Academyでなければならないのか】', level=2)

add_body(
    '「AI人材育成」を掲げるサービスは数多く存在します。'
    'Aidemy、キカガク、SIGNATE、Reskilling Camp、ブレインパッド、DataRobot、Schoo、Udemyなど、'
    '主要8社を含む市場調査を実施しました。その結果、明らかになった事実があります。'
)

add_body(
    '「製造業特化」と「伴走型実践」を同時に満たすスクールは、市場にゼロでした。',
    bold=True,
    align=WD_ALIGN_PARAGRAPH.CENTER
)

add_empty()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('（※ここにAX Academyの差別化ポジション概念図を配置）')
run.font.size = Pt(9)
run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
run.italic = True

# 一般AIスクールとの違い
add_empty()
p = doc.add_paragraph()
run = p.add_run('〈一般的なAIスクールとの違い〉')
run.bold = True
run.font.size = Pt(11)

add_body(
    'Aidemy、キカガク、SIGNATE、Schoo、Udemyなどの一般的なAIスクールは、'
    'プログラミングやデータサイエンスの汎用的なスキルを幅広く教えることを目的としています。'
    'カリキュラムは動画教材や座学が中心であり、用意されたケーススタディで学ぶスタイルが一般的です。'
    '修了後に得られるのは修了証や資格であり、現場での実装経験は含まれていません。'
)

add_body(
    'これらのスクールは、エンジニアやデータサイエンティストを目指す方には有効です。'
    'しかし、製造業の現場でAIを推進するリーダーを育てるという目的には適合しません。'
    '予知保全や品質管理といった製造業固有の課題に特化した内容が不足しており、'
    '自社のデータや業務課題を使った実践の機会もありません。'
)

add_body(
    'AX Academyは予知保全、品質管理、生産最適化、SCMという製造業の4領域に完全特化し、'
    '自社データを用いた実践により、受講期間中にPoC1件を現場で実装するところまで到達します。'
)

# コンサルとの違い
add_empty()
p = doc.add_paragraph()
run = p.add_run('〈製造業向けコンサルティングとの違い〉')
run.bold = True
run.font.size = Pt(11)

add_body(
    'ブレインパッドやDataRobotなどが提供する製造業向けのAIコンサルティングは、'
    '専門性の高い提案と実装支援を強みとしています。'
    'しかし、案件ごとに数百万円規模の費用が発生し、'
    'コンサルタントが去った後に社内にノウハウが残りにくいという構造的な課題があります。'
)

add_body(
    'AX Academyは、この外部依存の構造そのものを断ち切ることを目的としています。'
    'コンサルタントに代わるAX推進リーダーを社内に育成することで、'
    '2件目以降のAIプロジェクトは自社の力で推進できるようになります。'
)

# 4つの差別化軸
add_empty()
p = doc.add_paragraph()
run = p.add_run('〈AX Academyだけが持つ4つの差別化軸〉')
run.bold = True
run.font.size = Pt(11)

add_body(
    '1つ目は「製造業特化」です。予知保全、品質管理、生産最適化、SCMの4領域に絞り込み、'
    '製造業の現場が本当に必要とするAIスキルだけを届けます。'
)

add_body(
    '2つ目は「伴走型実践」です。動画を見て終わり、座学で知識を得て終わりではありません。'
    '2ヶ月間、メンターが受講者に寄り添い、自社の課題を素材にしながら実践的に学びを進めます。'
)

add_body(
    '3つ目は「フロー体験設計」です。チクセントミハイのフロー理論に基づき、'
    '難易度を段階的に調整しながら成功体験を積み上げていきます。'
    '「AIは難しい」という壁を、構造的に乗り越える仕組みです。'
)

add_body(
    '4つ目は「外部依存の断ち切り」です。研修やコンサルティングを繰り返す従来のモデルとは対極に、'
    '社内にAI推進の核となる人材を定着させます。'
    '1回の投資で、繰り返しの外部コストから解放される構造です。'
)

add_body(
    'これら4つの軸を同時に満たすサービスは、'
    '主要8社を含む市場調査の結果、AX Academy以外に確認されていません。',
    bold=True
)

add_separator()

# ========================================
# モニター募集概要
# ========================================
add_heading_styled('【第1期モニター募集概要】5名限定', level=2)

add_body('AX Academyでは、第1期モニター受講企業を5名限定で募集いたします。')

# テーブルで募集概要を作成
table = doc.add_table(rows=7, cols=2)
table.style = 'Light Grid Accent 1'
table.alignment = WD_TABLE_ALIGNMENT.CENTER

data = [
    ('募集人数', '5名限定（先着・選考制）'),
    ('対象', '年商50〜500億円規模の中堅製造業に勤務する\nDX推進担当者、経営者、製造部門責任者'),
    ('受講形式', '完全オンライン・伴走型（2ヶ月間）'),
    ('モニター特別価格', '60万円/名（税別）※通常価格100万円の40%OFF'),
    ('通常価格', '1名：100万円/名\n2名：87.5万円/名（12.5%割引）\n3名以上：70万円/名（30%割引）'),
    ('助成金', '人材開発支援助成金（最大75%補助）\n教育訓練給付金（最大70%補助）等に対応'),
    ('開講時期', '近日開講予定\n（お問い合わせから最短2週間でキックオフ可能）'),
]

for i, (label, value) in enumerate(data):
    cell_label = table.rows[i].cells[0]
    cell_value = table.rows[i].cells[1]

    run = cell_label.paragraphs[0].add_run(label)
    run.bold = True
    run.font.size = Pt(10)

    run = cell_value.paragraphs[0].add_run(value)
    run.font.size = Pt(10)

# 導入ステップ
add_empty()
p = doc.add_paragraph()
run = p.add_run('■ 導入までのステップ')
run.bold = True
run.font.size = Pt(11)

steps = [
    ('STEP 1：ヒアリング（30〜60分・無料）', '現状の課題・体制・目標をお伺いします。'),
    ('STEP 2：カスタム提案書の作成', '自社の状況に合わせた個別カリキュラム提案書をご提供いたします。'),
    ('STEP 3：ご契約', '内容にご納得いただいた上で正式契約となります。'),
    ('STEP 4：キックオフ（最短2週間）', '学習を開始いたします。'),
    ('STEP 5：2ヶ月間の伴走支援（PoC実装まで）', '修了後フォローアップも含めて完走をサポートいたします。'),
]

for title, desc in steps:
    p = doc.add_paragraph()
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(10.5)
    p2 = doc.add_paragraph(desc)
    for run in p2.runs:
        run.font.size = Pt(10.5)

# 成果
add_empty()
p = doc.add_paragraph()
run = p.add_run('■ 2ヶ月間で得られる成果')
run.bold = True
run.font.size = Pt(11)

add_bullet('PoC（概念実証）1件を現場で実装')
add_bullet('現場課題の明確化と優先順位の整理')
add_bullet('PoC実行計画の策定')
add_bullet('AI活用の推進ロードマップ策定')
add_bullet('経営報告用KPIの整備')

add_separator()

# ========================================
# 事業責任者コメント
# ========================================
add_heading_styled('【事業責任者コメント】', level=2)

add_empty()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('（※ここに事業責任者 熊切悠真の顔写真を配置）')
run.font.size = Pt(9)
run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
run.italic = True

add_body('株式会社TERRAISE CSO・AX Academy事業責任者　熊切悠真', bold=True)

comment_text = (
    '「日本の中堅製造業の皆さまは、AIに対して極めて真剣に向き合っておられます。'
    'しかし、その真剣さに応えられるサービスが、これまでほとんど存在しませんでした。\n\n'
    '汎用的な内容を教えて終わり。提案書を渡して終わり。'
    'そのたびにコストだけが積み上がり、スキルは社内に残らない。'
    '私はこの構造そのものに、強い課題意識を持っていました。\n\n'
    'AX Academyは、受講企業の現場でAIが実際に動くまで、責任を持って伴走します。'
    '座学で終わらせるつもりは一切ありません。'
    '2ヶ月後、担当者の方が自分の手でAIを動かし、現場の課題を解決している'
    '——その瞬間をつくるために、このスクールを立ち上げました。\n\n'
    '私たちが大切にしているのは、学びの中で『没頭』が生まれる体験設計です。'
    '心理学者チクセントミハイが提唱したフロー理論に基づき、'
    '適切な難易度の課題に集中し、手を動かしながら実感を伴って学べるプログラムを設計しています。'
    '知識を詰め込むのではなく、自らの手で成果を生み出した実感こそが、'
    '現場で使い続ける力になると確信しています。\n\n'
    '『研修・コンサルの繰り返しを、これで終わらせる』。これは私たちの覚悟の言葉です。'
    '外部に頼り続けるのではなく、自社の力でAIを活用し続けられる組織へ。'
    'AX Academyは、その転換点を共につくるパートナーでありたいと考えています。」'
)

p = doc.add_paragraph()
run = p.add_run(comment_text)
run.font.size = Pt(10.5)
run.italic = True
p.paragraph_format.left_indent = Cm(1.0)

add_separator()

# ========================================
# 今後の展望
# ========================================
add_heading_styled('【今後の展望】', level=2)

add_body(
    'AX Academyでは、以下のロードマップに基づき、'
    '日本の製造業におけるAI人材育成の基盤構築を推進してまいります。'
)

add_empty()
p = doc.add_paragraph()
run = p.add_run('短期（2026年内）：')
run.bold = True
run.font.size = Pt(10.5)
p2 = doc.add_paragraph(
    '第1期モニター5名の完遂と成果創出に全力を注ぐとともに、'
    '経済産業省「Reスキル講座」認定の取得を推進いたします。'
    '認定取得により、受講企業の費用負担をさらに軽減する体制を整えます。'
)

add_empty()
p = doc.add_paragraph()
run = p.add_run('中期（2027〜2028年）：')
run.bold = True
run.font.size = Pt(10.5)
p2 = doc.add_paragraph(
    '修了企業同士のネットワークを構築し、AIナレッジ共有コミュニティを形成します。'
    '業種・テーマ別の専門特化コースを追加開発するとともに、'
    '製造業で培った知見を活かし、金融・小売など他業種への横展開も計画しております。'
)

add_empty()
p = doc.add_paragraph()
run = p.add_run('長期ビジョン：')
run.bold = True
run.font.size = Pt(10.5)
p2 = doc.add_paragraph(
    '日本の中堅製造業が外部依存から脱却し、AIを自走活用できる組織へと変革すること。'
    'AX Academyは、その実現に向けた伴走者であり続けます。'
)

add_empty()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('（※ここにモニター募集概要カード＋QRコードを配置）')
run.font.size = Pt(9)
run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
run.italic = True

add_separator()

# ========================================
# 会社概要
# ========================================
add_heading_styled('【会社概要】', level=2)

table2 = doc.add_table(rows=7, cols=2)
table2.style = 'Light Grid Accent 1'
table2.alignment = WD_TABLE_ALIGNMENT.CENTER

company_data = [
    ('社名', '株式会社TERRAISE'),
    ('代表取締役', '舟橋'),
    ('CSO', '熊切悠真'),
    ('設立', '2025年1月31日'),
    ('所在地', '〒181-0012 東京都三鷹市上連雀1-12-17\n三鷹ビジネスパーク B1F'),
    ('事業内容', '製造業特化AIスクール「AX Academy」の運営、\nAI人材育成プログラムの企画・提供'),
    ('URL', 'https://terraise.studio.site/'),
]

for i, (label, value) in enumerate(company_data):
    cell_label = table2.rows[i].cells[0]
    cell_value = table2.rows[i].cells[1]

    run = cell_label.paragraphs[0].add_run(label)
    run.bold = True
    run.font.size = Pt(10)

    run = cell_value.paragraphs[0].add_run(value)
    run.font.size = Pt(10)

add_separator()

# ========================================
# お問い合わせ先
# ========================================
add_heading_styled('【本件に関するお問い合わせ先】', level=2)

add_body('株式会社TERRAISE AX Academy事業部', bold=True)
add_body('メール：ax@terraise.co.jp')
add_body('所在地：〒181-0012 東京都三鷹市上連雀1-12-17 三鷹ビジネスパーク B1F')
add_body('URL：https://terraise.studio.site/')

add_empty()
add_body(
    'まずはヒアリング（30〜60分・無料）からお気軽にご連絡ください。'
    '第1期モニターの募集枠は5名限定です。定員に達し次第、受付を終了いたします。',
    bold=True
)

add_empty()
add_body(
    '本プレスリリースに関する取材・掲載のご依頼も上記連絡先にてお受けしております。'
    '事業責任者へのインタビュー取材にも対応可能です。',
    italic=True
)

# 保存
output_path = '/home/user/-/AX_Academy_プレスリリース_PRTIMES入稿用.docx'
doc.save(output_path)
print(f'ドキュメント生成完了: {output_path}')
